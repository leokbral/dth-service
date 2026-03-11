from pathlib import Path

from docx import Document

from src.core.docx_processor import DocxProcessor


def test_convert_table_preserves_rowspan_and_colspan(tmp_path):
    doc_path = Path(tmp_path) / "table_merges.docx"

    doc = Document()
    table = doc.add_table(rows=3, cols=3)

    # Horizontal merge on first row: [A1 spans col 1-2] [A3]
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "A1 merged"
    table.cell(0, 2).text = "A3"

    # Vertical merge on third column: [A3 over B3]
    table.cell(0, 2).merge(table.cell(1, 2))
    table.cell(0, 2).text = "A3+B3"

    table.cell(1, 0).text = "B1"
    table.cell(1, 1).text = "B2"

    table.cell(2, 0).text = "C1"
    table.cell(2, 1).text = "C2"
    table.cell(2, 2).text = "C3"

    doc.save(doc_path)

    processor = DocxProcessor()
    html = processor.convert(str(doc_path))

    assert "<table" in html
    assert '<td colspan="2">A1 merged</td>' in html
    assert '<td rowspan="2">A3+B3</td>' in html
    assert "<td>B1</td>" in html
    assert "<td>B2</td>" in html
    assert "<td>C3</td>" in html
